"""Build the ML edition of the ETO data kit: guide workbook (catalog, dictionary, ML use cases, baselines, previews),
full CSV bundle, SQL schema."""
import os, zipfile
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from catalog import catalog, F

BASE = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(BASE, "csv")
ORDER = ["rfqs", "rfq_documents", "rfq_requirements", "contract_clauses", "quotes", "quote_lines", "supplier_quotes",
         "quote_outcomes", "projects", "project_bom_actuals", "customers", "product_families", "item_master", "suppliers",
         "competitors", "estimators", "labour_rates", "commodity_prices", "shop_capacity"]
data = {t: pd.read_csv(os.path.join(CSV, f"{t}.csv")) for t in ORDER + ["ml_training_table"]}
ml = data["ml_training_table"]
base = pd.read_csv(os.path.join(BASE, "ml_baseline_results.csv")).set_index("use_case")
closed = int((ml.y_project_status == "Closed").sum())
decided = int(ml.y_is_won.notna().sum())


def b(uc):
    if uc not in base.index:
        return "", "", ""
    r = base.loc[uc]
    return r.metric, r.model_score, r.naive_baseline


# (area, use case, ML task, target, key features, tables, rows, approach, business value, phase, notes) ; baseline key
UC = [
    ("Bid decision", "Bid / no-bid scoring", "Binary classification", "y_is_bid", "spec clarity, credit rating, LD cap, shop load next 12 weeks, complexity, bidders, delivery tightness, segment",
     "ml_training_table", len(ml), "Time split (train ≤2024, test ≥Jul-2025)", "Stop spending estimator time on unwinnable or risky RFQs", "Phase 2", ""),
    ("Bid decision", "Quote-effort prediction", "Regression", "q_quote_preparation_hrs", "complexity, document pages, family, requirements count",
     "ml_training_table", int(ml.q_quote_preparation_hrs.notna().sum()), "Random 75/25", "Plan bid-team capacity; feeds bid/no-bid", "Phase 2", ""),
    ("Win & price", "Win probability", "Binary classification", "y_is_won", "quoted price vs estimated cost, margin, lead-time gap, segment, bidders, customer win history",
     "ml_training_table", decided, "Time split", "Prioritise follow-up; basis for price optimisation", "Phase 2", ""),
    ("Win & price", "Price optimisation / price-to-win", "Win-probability curve over price", "y_is_won (with q_quoted_price_inr varied)", "Win model + cost estimate; simulate margin vs win probability",
     "ml_training_table", decided, "Uses win model; validate on held-out period", "Pick the price that maximises expected margin", "Phase 3", "No separate baseline – built on the win model"),
    ("Win & price", "Loss-reason prediction", "Multi-class classification", "y_loss_reason", "price ratio, lead-time gap, terms, competitor count",
     "ml_training_table", int((ml.y_outcome == "Lost").sum()), "Time split", "Flag likely loss reason before submission", "Phase 3", "Hard: competitor price unknown at quote time"),
    ("Win & price", "Negotiation discount", "Regression", "y_discount_pct", "target margin, bidders, segment, procurement style",
     "ml_training_table", int(ml.y_discount_pct.notna().sum()), "Random 75/25", "Build realistic negotiation headroom into price", "Phase 3", ""),
    ("Cost & hours", "Cost estimate correction (should-cost)", "Regression", "y_actual_total_cost_inr (via y_cost_overrun_ratio)", "estimator's cost build-up + specs + estimator experience + steel price",
     "ml_training_table", closed, "Random 75/25 on closed projects", "Halve estimating error; protect margin", "Phase 2", "Strongest use case in this data"),
    ("Cost & hours", "Cost prediction from specs only (no estimate)", "Regression", "y_actual_total_cost_inr", "capacity, weight, power, material, complexity, quantity, family",
     "ml_training_table", closed, "Random 75/25", "Instant budgetary quotes in hours, not weeks", "Phase 2", ""),
    ("Cost & hours", "Hours estimation – engineering", "Regression", "y_actual_engineering_hrs", "specs + estimator's hours + estimator experience",
     "ml_training_table", closed, "Random 75/25", "Better engineering resourcing and labour cost", "Phase 2", ""),
    ("Cost & hours", "Hours estimation – fabrication", "Regression", "y_actual_fabrication_hrs", "weight, complexity, material + estimator's hours",
     "ml_training_table", closed, "Random 75/25", "Shop loading and labour cost accuracy", "Phase 2", ""),
    ("Cost & hours", "Hours estimation – assembly", "Regression", "y_actual_assembly_hrs", "fabrication scope, complexity, quantity + estimator's hours",
     "ml_training_table", closed, "Random 75/25", "As above", "Phase 2", ""),
    ("Cost & hours", "Bought-out price prediction", "Regression", "supplier_quotes.quoted_unit_price_inr", "item category, standard cost, commodity driver, supplier quality and country, date",
     "supplier_quotes + item_master + suppliers", len(data["supplier_quotes"]), "Random 75/25", "Price bought-outs without waiting for vendor quotes", "Phase 1-2", ""),
    ("Cost & hours", "Draft BOM recommendation", "Multi-label classification / nearest neighbours", "quote_lines (item categories and sizes per quote)", "family, capacity, material, environment, power",
     "quote_lines + rfqs", len(data["quote_lines"]), "Group split by quote", "Start every estimate from a predicted BOM", "Phase 2", "No baseline run; item master is small (67 items) – extend with your own"),
    ("Cost & hours", "Similar past project retrieval", "Nearest neighbours / embeddings", "q_similar_past_quote_ref (weak label)", "family, capacity, material, design code, environment",
     "ml_training_table", int(ml.q_similar_past_quote_ref.notna().sum()), "Recall@k against the reference", "Reuse 60–80% of a past design and its costing", "Phase 1", ""),
    ("Execution risk", "Cost-overrun risk (>5%)", "Binary classification", "y_is_cost_overrun_gt5pct", "complexity, customisation, estimator experience, spec clarity, contingency",
     "ml_training_table", closed, "Random 75/25", "Add contingency where it is really needed", "Phase 2", ""),
    ("Execution risk", "Margin-erosion prediction", "Regression", "y_margin_erosion_pts", "as above + terms",
     "ml_training_table", closed, "Random 75/25", "Margin protection in approval workflow", "Phase 3", ""),
    ("Execution risk", "Delivery-delay prediction", "Regression", "y_delivery_slip_weeks", "shop load next 12 weeks, complexity, customisation, spec clarity, lead-time gap",
     "ml_training_table", closed, "Random 75/25", "Promise delivery dates the plant can keep", "Phase 3", ""),
    ("Execution risk", "LD (penalty) exposure", "Binary classification", "y_ld_paid_flag", "LD terms + delay drivers",
     "ml_training_table", closed, "Random 75/25", "Price or negotiate LD risk before signing", "Phase 3", ""),
    ("Execution risk", "Engineering-change count", "Count regression", "y_num_engineering_changes", "complexity, customisation, spec clarity",
     "ml_training_table", closed, "Random 75/25", "Plan engineering buffer", "Phase 3", ""),
    ("Execution risk", "Missed-BOM-item risk", "Binary classification", "y_n_items_added_after_quote > 0", "complexity, customisation, spec clarity",
     "ml_training_table + project_bom_actuals", closed, "Random 75/25", "Flag quotes likely to miss scope", "Phase 3", ""),
    ("Documents & NLP", "Requirement categorisation (NLP)", "Text classification", "rfq_requirements.category", "requirement_text",
     "rfq_requirements", len(data["rfq_requirements"]), "Random 75/25", "Auto-build compliance matrices from specs", "Phase 1", ""),
    ("Documents & NLP", "Compliance / deviation suggestion", "Text + context classification", "rfq_requirements.compliance_status", "requirement text + RFQ terms and clarity",
     "rfq_requirements + rfqs", len(data["rfq_requirements"]), "Random 75/25", "Draft deviations and clarifications automatically", "Phase 1-2", ""),
    ("Documents & NLP", "Contract clause risk scoring", "Classification (rules + ML)", "contract_clauses.risk_level", "clause_type, customer_position vs our_standard_position",
     "contract_clauses", len(data["contract_clauses"]), "Random 75/25", "Highlight high-risk T&Cs for commercial review", "Phase 1", "Labels are rule-based here – near-perfect accuracy expected; test the pipeline, not the model"),
    ("Documents & NLP", "RFQ PDF field and clause extraction (LLM)", "Information extraction", "RFQ_PDF_Answer_Key.xlsx", "The 3 customer PDFs",
     "customer_rfq_pdfs", 3, "Field-level precision / recall vs answer key", "Cut RFQ reading time from days to minutes", "Phase 1", "Only 3 documents – an evaluation set, not a training set"),
    ("Planning & forecasting", "Commodity price forecasting", "Time series", "commodity_prices.price", "price history per commodity",
     "commodity_prices", len(data["commodity_prices"]), "Rolling-origin backtest", "Set quote validity and escalation clauses", "Phase 3", "Short monthly series – replace with a real index feed"),
    ("Planning & forecasting", "Shop-load forecasting", "Time series", "shop_capacity.utilisation_pct", "work centre history, order book (won projects)",
     "shop_capacity + projects", len(data["shop_capacity"]), "Rolling-origin backtest", "Feeds lead-time promise and bid/no-bid", "Phase 3", ""),
    ("Analytics", "Estimator bias analysis", "Descriptive / mixed-effects model", "y_cost_overrun_ratio by q_estimator_id", "estimator, experience, family, complexity",
     "ml_training_table", closed, "n/a", "Targeted coaching and estimate review rules", "Phase 1", "Bias is built into the generator (less experience = more under-estimation)"),
    ("Analytics", "Customer segmentation & value", "Clustering", "customer-level aggregates", "win rate, margin, payment days, volume, terms",
     "customers + ml_training_table", len(data["customers"]), "n/a", "Focus sales effort on profitable accounts", "Phase 3", ""),
]

ARIAL = "Arial"
HF, HB = Font(name=ARIAL, bold=True, color="FFFFFF", size=10), PatternFill("solid", start_color="1F3864")
BODY, BOLD = Font(name=ARIAL, size=10), Font(name=ARIAL, size=10, bold=True)
WRAP, THIN = Alignment(wrap_text=True, vertical="top"), Side(style="thin", color="D9D9D9")
PRIO = {"Must": "F8CBAD", "Should": "FFE699", "Nice to have": "E2EFDA"}
wb = Workbook()


def header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c)
        x.font, x.fill, x.alignment = HF, HB, Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, w):
    for j, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = x


def put(ws, r, c, v, font=BODY, wrap=True, fmt=None):
    x = ws.cell(row=r, column=c, value=None if (isinstance(v, float) and np.isnan(v)) else v)
    x.font = font
    x.alignment = WRAP if wrap else Alignment(vertical="top")
    if fmt:
        x.number_format = fmt
    return x


# ------------------------------------------------ README
ws = wb.active
ws.title = "README"
txt = [
    ("ETO AI Quoting – ML Dataset Kit", Font(name=ARIAL, size=16, bold=True, color="1F3864")),
    ("Synthetic but internally consistent data for a generic engineer-to-order manufacturer, Jan-2021 to Sep-2026, values in INR.", Font(name=ARIAL, size=10, italic=True)),
    ("", BODY),
    ("Start here", BOLD),
    ("1. ml_training_table.csv (in the zip) – one row per RFQ, 109 columns, ready for modelling. Column prefixes:", BODY),
    ("     f_ = known when the RFQ arrives;  q_ = our quote (estimate, price, lead time);  y_ = outcomes known later (targets).", BODY),
    ("     Never use y_ columns as features for quote-time models. Use the 'split' column for a time-based train / valid / test split.", BODY),
    ("2. ML_Use_Cases – 28 use cases with target, features, rows available and a quick baseline result where one was run.", BODY),
    ("3. The 19 relational CSVs hold the full detail (requirements text, quote lines, supplier quotes, BOM actuals, capacity, commodities).", BODY),
    ("", BODY),
    ("Tabs", BOLD),
    ("ML_Use_Cases · Data_Catalog · Field_Dictionary · Relationships · Row_Counts · then a 300-row preview of every table (full data in the zip).", BODY),
    ("", BODY),
    ("Signal built into the data (so models have something real to learn)", BOLD),
    ("• Cost and hours scale with size, power, material, complexity and customisation; estimators under-estimate complex, custom jobs and less", BODY),
    ("  experienced estimators under-estimate more.  • Win probability falls with price above competitors, longer lead times and more bidders.", BODY),
    ("• No-bid is driven by spec clarity, credit risk, harsh LDs, shop load and tight delivery.  • Delays follow shop load, complexity and unclear specs,", BODY),
    ("  and trigger LDs.  • Supplier prices follow commodity indices, supplier country and quality; supplier lead times follow supplier history and load.", BODY),
    ("• Compliance responses depend on clause type and RFQ terms.  • Customer-history features in the ML table use only information known before each RFQ.", BODY),
    ("", BODY),
    ("Limitations", BOLD),
    ("Synthetic: models trained here will not transfer to your products. Use it to build pipelines, features and evaluation; retrain on your own history.", BODY),
    ("Requirement text is template-based (NLP scores are optimistic). Contract risk labels are rule-based. Item master is small (67 items).", BODY),
    ("RFQ-2023-0105, RFQ-2023-0094 and RFQ-2026-0563 match the three customer PDFs; their requirements carry real clause and page references.", BODY),
    ("Regenerate: generate_eto_data.py → align_rfq_pdfs.py → build_ml_table.py → ml_baselines.py → build_ml_kit.py (all in the zip).", BODY),
]
for i, (t, f_) in enumerate(txt, 1):
    ws.cell(row=i, column=1, value=t).font = f_
widths(ws, [150])

# ------------------------------------------------ ML_Use_Cases
ws = wb.create_sheet("ML_Use_Cases")
cols = ["#", "Area", "Use case", "ML task", "Target", "Key features", "Tables", "Rows available", "Evaluation approach",
        "Baseline metric", "Model score", "Naive baseline", "Business value", "Phase", "Notes"]
header(ws, 1, cols)
for i, u in enumerate(UC, 2):
    area, name, task, tgt, feats, tabs, rows, appr, value, phase, note = u
    metric, score, naive = b(name)
    if name in base.index and base.loc[name].note and not note:
        note = base.loc[name].note
    vals = [i - 1, area, name, task, tgt, feats, tabs, int(rows), appr, metric, score if score != "" else None, naive if naive != "" else None, value, phase, note]
    for j, v in enumerate(vals, 1):
        put(ws, i, j, v, fmt="#,##0" if j == 8 else ("0.000" if j in (11, 12) else None))
        ws.cell(row=i, column=j).border = Border(bottom=THIN)
widths(ws, [4, 16, 30, 20, 26, 40, 22, 10, 22, 20, 9, 9, 34, 9, 40])
ws.auto_filter.ref = f"A1:O{len(UC) + 1}"
n = len(UC) + 3
put(ws, n, 1, "Model score vs naive baseline: AUC – higher is better, naive 0.5; MAPE / MAE – lower is better, naive is the estimator's own figure or the mean; "
              "F1 / accuracy – higher is better, naive is the majority class. Untuned gradient boosting (scikit-learn) – indicative only.",
    Font(name=ARIAL, size=9, italic=True), wrap=False)

# ------------------------------------------------ Data_Catalog
ws = wb.create_sheet("Data_Catalog")
cat = catalog.copy()
mlrow = {"table_name": "ml_training_table", "in_sample_data": "Yes", "domain": "ML (derived)",
         "what_it_contains": "One row per RFQ joining RFQ, customer, quote, outcome and project actuals; f_/q_/y_ column groups and a time split",
         "ai_use_cases_fed": "All tabular ML use cases", "typical_source_system": "Built from the tables above (build_ml_table.py)",
         "grain": "1 row per RFQ", "refresh_frequency": "On rebuild", "minimum_history_needed": "n/a", "priority": "Must",
         "data_owner": "Data / ML team", "notes": "Feature store starting point"}
cat = pd.concat([pd.DataFrame([mlrow]), cat], ignore_index=True)
cols = ["table_name", "in_sample_data", "domain", "what_it_contains", "ai_use_cases_fed", "typical_source_system", "grain",
        "refresh_frequency", "minimum_history_needed", "priority", "data_owner", "notes"]
header(ws, 1, [c.replace("_", " ").title() for c in cols] + ["Rows in sample"])
for i, row in enumerate(cat.to_dict("records"), 2):
    for j, c in enumerate(cols, 1):
        x = put(ws, i, j, row[c])
        x.border = Border(bottom=THIN)
        if c == "priority":
            x.fill = PatternFill("solid", start_color=PRIO[row[c]])
    put(ws, i, len(cols) + 1, len(data[row["table_name"]]) if row["table_name"] in data else "n/a", fmt="#,##0")
widths(ws, [22, 10, 15, 40, 40, 34, 22, 14, 18, 11, 20, 36, 12])

# ------------------------------------------------ Field_Dictionary
ws = wb.create_sheet("Field_Dictionary")
header(ws, 1, ["Table", "Field", "Description", "Key", "Data type", "Example value", "% populated"])


def dtype(s):
    if pd.api.types.is_integer_dtype(s):
        return "INTEGER"
    if pd.api.types.is_float_dtype(s):
        return "NUMBER"
    nn = s.dropna().astype(str)
    return "DATE" if len(nn) and nn.str.match(r"^\d{4}-\d{2}-\d{2}$").all() else "VARCHAR"


def ml_desc(c):
    if c in ("rfq_id", "quote_id", "project_id", "customer_id"):
        return "Key", "PK" if c == "rfq_id" else "FK"
    if c == "split":
        return "Time-based split: train (RFQ before 2025), valid (Jan–Jun 2025), test (from Jul 2025)", ""
    if c == "received_date":
        return "RFQ received date", ""
    pre = {"f_": "Feature known at RFQ time", "q_": "Our quote / estimate", "y_": "Outcome or actual (target – known later)"}[c[:2]]
    return f"{pre}: {c[2:].replace('_', ' ')}", ""


r = 2
for t in ["ml_training_table"] + ORDER:
    df = data[t]
    desc = {f_: (d, k) for f_, d, k in F.get(t, [])}
    for col in df.columns:
        d, k = desc.get(col, ml_desc(col) if t == "ml_training_table" else ("", ""))
        s = df[col]
        ex = s.dropna().iloc[0] if s.notna().any() else ""
        for j, v in enumerate([t, col, d, k, dtype(s), str(ex)[:60], round(float(s.notna().mean()), 3)], 1):
            put(ws, r, j, v, wrap=(j == 3), fmt="0%" if j == 7 else None)
        r += 1
widths(ws, [20, 34, 56, 10, 11, 34, 11])
ws.auto_filter.ref = f"A1:G{r - 1}"

# ------------------------------------------------ Relationships & Row_Counts
ws = wb.create_sheet("Relationships")
header(ws, 1, ["Child table", "Child column", "Parent table", "Parent column", "Cardinality"])
REL = [("rfqs", "customer_id", "customers", "customer_id", "many-to-one"), ("rfqs", "family_id", "product_families", "family_id", "many-to-one"),
       ("rfq_documents", "rfq_id", "rfqs", "rfq_id", "many-to-one"), ("rfq_requirements", "rfq_id", "rfqs", "rfq_id", "many-to-one"),
       ("contract_clauses", "rfq_id", "rfqs", "rfq_id", "many-to-one"), ("quotes", "rfq_id", "rfqs", "rfq_id", "one-to-one (bids only)"),
       ("quotes", "estimator_id", "estimators", "estimator_id", "many-to-one"), ("quotes", "similar_past_quote_ref", "quotes", "quote_id", "many-to-one (self)"),
       ("quote_lines", "quote_id", "quotes", "quote_id", "many-to-one"), ("quote_lines", "item_id", "item_master", "item_id", "many-to-one (material lines)"),
       ("supplier_quotes", "quote_id", "quotes", "quote_id", "many-to-one"), ("supplier_quotes", "supplier_id", "suppliers", "supplier_id", "many-to-one"),
       ("supplier_quotes", "item_id", "item_master", "item_id", "many-to-one"), ("quote_outcomes", "rfq_id", "rfqs", "rfq_id", "one-to-one"),
       ("quote_outcomes", "quote_id", "quotes", "quote_id", "one-to-one (blank for no-bid)"), ("quote_outcomes", "winning_competitor_id", "competitors", "competitor_id", "many-to-one"),
       ("projects", "quote_id", "quotes", "quote_id", "one-to-one (won quotes)"), ("projects", "customer_id", "customers", "customer_id", "many-to-one"),
       ("project_bom_actuals", "project_id", "projects", "project_id", "many-to-one"), ("project_bom_actuals", "item_id", "item_master", "item_id", "many-to-one"),
       ("item_master", "preferred_supplier_id", "suppliers", "supplier_id", "many-to-one"), ("item_master", "commodity_driver", "commodity_prices", "commodity", "many-to-one"),
       ("ml_training_table", "rfq_id / quote_id / project_id", "rfqs / quotes / projects", "keys", "one row per RFQ (derived)")]
for i, row in enumerate(REL, 2):
    for j, v in enumerate(row, 1):
        put(ws, i, j, v, wrap=False)
widths(ws, [22, 30, 24, 18, 30])

ws = wb.create_sheet("Row_Counts")
header(ws, 1, ["Table", "Rows", "Columns", "Date range"])
for i, t in enumerate(["ml_training_table"] + ORDER, 2):
    df = data[t]
    dcol = next((c for c in ["received_date", "quote_date", "order_date", "week_start", "month", "quote_received_date"] if c in df.columns), None)
    rng = f"{df[dcol].min()} to {df[dcol].max()}" if dcol else ""
    for j, v in enumerate([t, len(df), df.shape[1], rng], 1):
        put(ws, i, j, v, wrap=False, fmt="#,##0" if j in (2, 3) else None)
widths(ws, [24, 12, 10, 30])

# ------------------------------------------------ previews
for t in ["ml_training_table"] + ORDER:
    df = data[t].head(300)
    ws = wb.create_sheet(("ml_table_preview" if t == "ml_training_table" else t)[:31])
    header(ws, 1, list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for j, col in enumerate(df.columns, 1):
        L = get_column_letter(j)
        ws.column_dimensions[L].width = min(max(10, len(col) + 2, int(df[col].fillna("").astype(str).str.len().quantile(.9)) + 2), 50)
        fmt = "#,##0" if col.endswith("_inr") else None
        for cell in ws[L][1:]:
            cell.font = BODY
            if fmt:
                cell.number_format = fmt
    ws.auto_filter.ref = ws.dimensions

out = os.path.join(BASE, "ETO_AI_Quoting_ML_Dataset_Guide.xlsx")
wb.save(out)

# ------------------------------------------------ DDL
TYPE = {"INTEGER": "INTEGER", "NUMBER": "NUMERIC(18,4)", "DATE": "DATE", "VARCHAR": "VARCHAR(1000)"}
PK = {"customers": ["customer_id"], "product_families": ["family_id"], "item_master": ["item_id"], "suppliers": ["supplier_id"],
      "competitors": ["competitor_id"], "estimators": ["estimator_id"], "labour_rates": ["resource_type", "effective_from"],
      "commodity_prices": ["month", "commodity"], "rfqs": ["rfq_id"], "rfq_documents": ["document_id"], "rfq_requirements": ["requirement_id"],
      "contract_clauses": ["rfq_id", "clause_type"], "quotes": ["quote_id"], "quote_lines": ["quote_id", "line_no"],
      "supplier_quotes": ["supplier_quote_id"], "quote_outcomes": ["rfq_id"], "projects": ["project_id"], "project_bom_actuals": None,
      "shop_capacity": ["week_start", "work_center_id"], "ml_training_table": ["rfq_id"]}
ddl = ["-- ETO AI Quoting – schema for the ML dataset kit (ANSI SQL; adjust types for your platform)", ""]
master = ["customers", "product_families", "suppliers", "item_master", "competitors", "estimators", "labour_rates", "commodity_prices"]
for t in master + [x for x in ORDER if x not in master] + ["ml_training_table"]:
    df = data[t]
    desc = {f_: d for f_, d, _ in F.get(t, [])}
    items = [(f"{c} {TYPE[dtype(df[c])]}", desc.get(c, ml_desc(c)[0] if t == "ml_training_table" else "")) for c in df.columns]
    if PK[t]:
        items.append((f"PRIMARY KEY ({', '.join(PK[t])})", ""))
    body = [(f"    {(sql + (',' if k < len(items) - 1 else '')):<52}" + (f" -- {cm}" if cm else "")).rstrip() for k, (sql, cm) in enumerate(items)]
    ddl += [f"CREATE TABLE {t} ("] + body + [");", ""]
open(os.path.join(BASE, "eto_schema.sql"), "w").write("\n".join(ddl))

# ------------------------------------------------ zip
zp = os.path.join(BASE, "ETO_ML_dataset_full.zip")
with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    z.write(os.path.join(CSV, "ml_training_table.csv"), "eto_ml_dataset/ml_training_table.csv")
    for t in ORDER:
        z.write(os.path.join(CSV, f"{t}.csv"), f"eto_ml_dataset/relational/{t}.csv")
    z.write(os.path.join(BASE, "ml_baseline_results.csv"), "eto_ml_dataset/ml_baseline_results.csv")
    z.write(os.path.join(BASE, "eto_schema.sql"), "eto_ml_dataset/eto_schema.sql")
    for f_ in ["generate_eto_data.py", "rfq_pdf_content.py", "build_rfq_pdfs.py", "align_rfq_pdfs.py", "build_ml_table.py", "ml_baselines.py",
               "catalog.py", "build_ml_kit.py"]:
        z.write(os.path.join(BASE, f_), f"eto_ml_dataset/scripts/{f_}")
    pdir = os.path.join(BASE, "customer_rfq_pdfs")
    for f_ in sorted(os.listdir(pdir)):
        z.write(os.path.join(pdir, f_), f"eto_ml_dataset/customer_rfq_pdfs/{f_}")
print(out, os.path.getsize(out) // 1024, "KB;", zp, os.path.getsize(zp) // 1024, "KB")
