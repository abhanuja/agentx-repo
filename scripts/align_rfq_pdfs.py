"""Make the sample dataset match the three customer RFQ PDFs and write an extraction answer key.

Run after generate_eto_data.py and build_rfq_pdfs.py (regenerating the data overwrites these edits).
"""
import os, re
import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from rfq_pdf_content import RFQS

BASE = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(BASE, "csv")
PDFS = os.path.join(BASE, "customer_rfq_pdfs")
rd = lambda n: pd.read_csv(os.path.join(CSV, f"{n}.csv"))
wr = lambda df, n: df.to_csv(os.path.join(CSV, f"{n}.csv"), index=False)

rfqs, customers = rd("rfqs"), rd("customers")
reqs, clauses, docs = rd("rfq_requirements"), rd("contract_clauses"), rd("rfq_documents")

HEADER = {  # rfq-level fields as stated in each PDF
    "RFQ-2023-0105": dict(rfq_title="Reagent Dosing & Transfer Skid - 61.1 m3/h", end_user_industry="Mining", design_code="ASME VIII Div 1",
                          primary_material="SS316L", operating_environment="Indoor dusty", currency="INR", inco_terms="FOR site"),
    "RFQ-2023-0094": dict(rfq_title="Rail-Mounted Goliath Cranes 82.4 t SWL (2 nos)", design_code="FEM 1.001", primary_material="S355 / IS 2062 E350",
                          operating_environment="Coastal / saline", currency="USD", inco_terms="CIF"),
    "RFQ-2026-0563": dict(rfq_title="End-Cap Welding & Leak-Test Machines - 663 units/h (2 nos)", design_code="CE Machinery Directive",
                          primary_material="SS316L", operating_environment="Corrosive", currency="EUR", inco_terms="FOB"),
}
CUSTOMER_FIX = {"C0033": dict(industry="Mining"), "C0137": dict(customer_name="Delta Wassertechnik GmbH")}
ACCEPTED = {"RFQ-2023-0105": {"Warranty": "N"}, "RFQ-2023-0094": {"Liquidated damages": "N", "Payment terms": "N", "Price escalation": "N"},
            "RFQ-2026-0563": {"Liquidated damages": "N", "Performance guarantee": "N"}}


def norm(t):
    return re.sub(r"\s+", " ", t).strip()


# ---------------------------------------------------------------- customers & rfqs
for cid, fix in CUSTOMER_FIX.items():
    for k, v in fix.items():
        customers.loc[customers.customer_id == cid, k] = v
rfqs.loc[rfqs.customer_id == "C0033", "end_user_industry"] = "Mining"

if "clause_ref" not in reqs.columns:
    reqs.insert(3, "clause_ref", None)

key_req, key_hdr, new_docs = [], [], []
for r in RFQS:
    rid = r["rfq_id"]
    pdf_path = os.path.join(PDFS, r["file"])
    with pdfplumber.open(pdf_path) as pdf:
        pages = [norm(p.extract_text() or "") for p in pdf.pages]
        n_pages = len(pdf.pages)
    for k, v in HEADER[rid].items():
        rfqs.loc[rfqs.rfq_id == rid, k] = v
    rfqs.loc[rfqs.rfq_id == rid, "rfq_document_pages"] = n_pages
    row = rfqs[rfqs.rfq_id == rid].iloc[0]

    # requirements: replace with the clauses actually in the PDF
    allc = [(c, "3. Technical specification") for cls in r["clauses"].values() for c in cls] + [(r["makes_clause"], "4. Approved makes")] \
        + [(c, "5. Inspection and testing") for c in r["testing"]] + [(c, "6. Documentation") for c in r["docs"]] \
        + [(c, "7. Commercial terms and conditions") for c in r["commercial"]]
    reqs = reqs[reqs.rfq_id != rid]
    start = int(reqs.requirement_id.str[3:].astype(int).max()) + 1
    rows = []
    for i, ((no, cat, text, mand, status, dev, cost), section) in enumerate(allc):
        probe = norm(text)[:45]
        page = next((p + 1 for p, t in enumerate(pages) if probe in t), None)
        rows.append({"requirement_id": f"REQ{start + i:07d}", "rfq_id": rid, "category": cat, "clause_ref": no, "requirement_text": text,
                     "source_document": r["file"], "source_page": page, "mandatory": mand, "compliance_status": status,
                     "deviation_or_clarification": dev, "cost_impact": cost})
        key_req.append({"rfq_id": rid, "pdf_file": r["file"], "section": section, "clause_ref": no, "page": page, "category": cat,
                        "mandatory": mand, "requirement_text": text, "our_compliance_status": status, "our_deviation_or_clarification": dev})
    reqs = pd.concat([reqs, pd.DataFrame(rows)], ignore_index=True)

    for ct, val in ACCEPTED[rid].items():
        clauses.loc[(clauses.rfq_id == rid) & (clauses.clause_type == ct), "accepted_as_is"] = val
    clauses.loc[(clauses.rfq_id == rid) & (clauses.accepted_as_is.isna()), "accepted_as_is"] = "Y"

    if not (docs.file_name == r["file"]).any():
        new_docs.append({"document_id": None, "rfq_id": rid, "document_type": "RFQ / tender document (customer PDF)", "file_format": "pdf",
                         "file_name": r["file"], "pages": n_pages, "revision": "R0", "received_date": row.received_date,
                         "storage_path": f"/rfq_archive/{row.received_date[:4]}/{rid}/"})

    hdr = [("Customer", r["customer"]), ("Customer enquiry no.", r["customer_ref"]), ("Date of issue", r["issue_date"]),
           ("Bid due", r["due"]), ("Pre-bid query deadline", r["query_deadline"]), ("Bid validity", r["validity"]),
           ("Project", r["project"]), ("Title", r["title"]), ("Product family", row.family_id), ("Quantity", int(row.quantity)),
           ("Capacity", f"{row.capacity_value} {row.capacity_unit}"), ("Design code", row.design_code), ("Primary material", row.primary_material),
           ("Operating environment", row.operating_environment), ("Required delivery (weeks)", int(row.required_delivery_weeks)),
           ("Incoterm", row.inco_terms), ("Currency", row.currency), ("LD % per week", row.ld_clause_pct_per_week), ("LD cap %", row.ld_cap_pct),
           ("Warranty (months)", int(row.warranty_months)), ("Payment terms", row.payment_terms), ("Performance guarantee", row.performance_guarantee),
           ("Installation supervision in scope", row.scope_includes_installation), ("Civil works in scope", row.scope_includes_civil)]
    key_hdr += [{"rfq_id": rid, "pdf_file": r["file"], "field": f, "expected_value": v} for f, v in hdr]

if new_docs:
    nd = pd.DataFrame(new_docs)
    start = int(docs.document_id.str[1:].astype(int).max()) + 1
    nd["document_id"] = [f"D{start + i:06d}" for i in range(len(nd))]
    docs = pd.concat([docs, nd], ignore_index=True)

reqs["source_page"] = reqs.source_page.astype("Int64")
for df, n in [(rfqs, "rfqs"), (customers, "customers"), (reqs, "rfq_requirements"), (clauses, "contract_clauses"), (docs, "rfq_documents")]:
    wr(df, n)

# ---------------------------------------------------------------- answer key workbook
wb = Workbook()
HF, HB, BF = Font(name="Arial", bold=True, color="FFFFFF", size=10), PatternFill("solid", start_color="1F3864"), Font(name="Arial", size=10)
ws = wb.active
ws.title = "README"
for i, t in enumerate([
    "RFQ PDF extraction – answer key",
    "Expected values for the three fictitious customer RFQ PDFs. Use it to score an LLM / document-AI extraction pipeline.",
    "Header_Fields: the RFQ-level fields an extractor should return per PDF (compare with the rfqs tab of the starter kit).",
    "Requirements: every numbered requirement clause in each PDF, with page number and category (matches rfq_requirements).",
    "Columns starting with 'our_' are the bidder's response (compliance and deviations), not text in the PDF. Use them to test a",
    "deviation-suggestion model; exclude them when scoring extraction.",
], 1):
    ws.cell(row=i, column=1, value=t).font = Font(name="Arial", size=14 if i == 1 else 10, bold=i == 1)
ws.column_dimensions["A"].width = 120
for name, rows in [("Header_Fields", key_hdr), ("Requirements", key_req)]:
    df = pd.DataFrame(rows)
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill = HF, HB
    for rec in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in rec])
    for j, col in enumerate(df.columns, 1):
        w = min(70, max(10, int(df[col].astype(str).str.len().quantile(.9)) + 2, len(col) + 2))
        ws.column_dimensions[get_column_letter(j)].width = w
        for c in ws[get_column_letter(j)][1:]:
            c.font = BF
            c.alignment = Alignment(wrap_text=w >= 60, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
wb.save(os.path.join(PDFS, "RFQ_PDF_Answer_Key.xlsx"))
print("requirements per PDF:", pd.DataFrame(key_req).groupby("rfq_id").size().to_dict(),
      "| missing page refs:", int(pd.DataFrame(key_req).page.isna().sum()))
